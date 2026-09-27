"""Word (.docx) export for Innovation Lab agent runs.

Renders the structured run result (summary, workflow trace, claims, sections,
findings, citations, suggestions) into a styled Word document via python-docx,
mirroring Eureka's "export the report" step for every agent.
"""

from __future__ import annotations

import os
import re
from datetime import datetime
from io import BytesIO
from typing import Any

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ACCENT = RGBColor(0x05, 0x96, 0x69)
DARK = RGBColor(0x06, 0x3F, 0x36)
GREY = RGBColor(0x6B, 0x72, 0x80)
BORDER = "059669"


def _shade(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _cell_text(cell, text: str, bold: bool = False, size: int = 9, color: RGBColor = DARK) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = color


def build_agent_docx(
    agent_slug: str,
    agent_label: str,
    phase: str,
    result: dict[str, Any],
) -> BytesIO:
    """Render a run result into a Word document buffer."""
    doc = Document()

    for section in doc.sections:
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(10.5)

    # --- Header block ----------------------------------------------------- #
    head = doc.add_heading("", level=0)
    run = head.add_run("IP-SAKTI Innovation Lab")
    run.font.color.rgb = ACCENT
    run.bold = True
    run.font.size = Pt(20)

    title = doc.add_heading("", level=1)
    tr = title.add_run(agent_label)
    tr.font.color.rgb = DARK
    sub = doc.add_paragraph()
    sr = sub.add_run(f"{agent_slug}  ·  {phase.replace('_', ' ')}  ·  {datetime.now().strftime('%d %b %Y, %H:%M')}")
    sr.font.size = Pt(9)
    sr.font.color.rgb = GREY
    doc.add_paragraph("—" * 78)

    # --- Summary / note --------------------------------------------------- #
    summary = str(result.get("summary") or "")
    if summary:
        p = doc.add_paragraph()
        r = p.add_run(summary)
        r.bold = True
        r.font.color.rgb = DARK
        r.font.size = Pt(11.5)
    note = str(result.get("note") or "")
    if note:
        p = doc.add_paragraph()
        r = p.add_run(note)
        r.italic = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = GREY

    # --- Execution trace -------------------------------------------------- #
    workflow = result.get("workflow") or []
    if workflow:
        doc.add_heading("Execution trace", level=2)
        for w in workflow:
            if isinstance(w, dict):
                label = w.get("label") or ""
                desc = w.get("description") or ""
            else:
                label, desc = str(w), ""
            p = doc.add_paragraph(style="List Number")
            r = p.add_run(label or "step")
            r.bold = True
            if desc:
                dr = p.add_run(" — " + desc)
                dr.font.size = Pt(9)
                dr.font.color.rgb = GREY

    # --- Claims ----------------------------------------------------------- #
    claims = result.get("claims") or []
    if claims:
        doc.add_heading("Draft claims", level=2)
        for i, c in enumerate(claims, start=1):
            p = doc.add_paragraph(style="List Number")
            p.add_run(f"Claim {i}").bold = True
            text = str(c).strip()
            if re.match(r"^\d+\.\s", text):
                text = re.sub(r"^\d+\.\s*", "", text)
            p.add_run(": " + text)

    # --- Sections (tables) ------------------------------------------------ #
    for sec in (result.get("sections") or []):
        title_text = str(sec.get("title") or "Section")
        doc.add_heading(title_text, level=2)
        caption = sec.get("caption")
        if caption:
            cp = doc.add_paragraph()
            cr = cp.add_run(str(caption))
            cr.italic = True
            cr.font.size = Pt(8.5)
            cr.font.color.rgb = GREY

        rows = sec.get("rows") or []
        if not rows:
            continue
        columns = sec.get("columns") or list(rows[0].keys())
        table = doc.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER

        hdr = table.rows[0].cells
        for j, col in enumerate(columns):
            _shade(hdr[j], "E7F6F0")
            _cell_text(hdr[j], str(col).replace("_", " ").title(), bold=True, color=ACCENT)

        for row in rows:
            cells = table.add_row().cells
            for j, col in enumerate(columns):
                _cell_text(cells[j], str(row.get(col, "") or ""))

        for ci, _col in enumerate(columns):
            table.columns[ci].width = Inches(max(0.9, 5.5 / max(len(columns), 1)))

        doc.add_paragraph()

    # --- Findings --------------------------------------------------------- #
    findings = result.get("findings") or []
    if findings:
        doc.add_heading("Findings", level=2)
        for f in findings:
            if isinstance(f, dict):
                title_f = f.get("title") or ""
                detail = f.get("detail") or ""
                sev = str(f.get("severity") or "info")
            else:
                title_f, detail, sev = str(f), "", "info"
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(title_f)
            r.bold = True
            r.font.color.rgb = DARK if sev != "error" else RGBColor(0xDC, 0x26, 0x26)
            if detail:
                p.add_run(" — " + str(detail))

    # --- Citations -------------------------------------------------------- #
    citations = result.get("citations") or []
    if citations:
        doc.add_heading(f"Citations ({len(citations)})", level=2)
        for c in citations:
            if isinstance(c, dict):
                act = c.get("act_title") or ""
                ref = c.get("section_reference") or ""
                auth = c.get("authority") or ""
                line = f"{act}" + (f" · {ref}" if ref else "") + (f" · {auth}" if auth else "")
            else:
                line = str(c)
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(line).font.size = Pt(9.5)

    # --- Suggestions ------------------------------------------------------ #
    suggestions = result.get("suggestions") or []
    if suggestions:
        doc.add_heading("Next steps", level=2)
        for s in suggestions:
            p = doc.add_paragraph(style="List Bullet")
            r = p.add_run(str(s))
            r.font.color.rgb = RGBColor(0xB4, 0x6A, 0x00)

    # --- Footer ----------------------------------------------------------- #
    doc.add_paragraph("—" * 78)
    foot = doc.add_paragraph()
    fr = foot.add_run(
        "AI-assisted research material — must be reviewed by a qualified professional. "
        "Not legal, medical, safety, regulatory, or patentability advice."
    )
    fr.italic = True
    fr.font.size = Pt(8)
    fr.font.color.rgb = GREY

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def export_path(slug: str) -> str:
    export_dir = os.path.join(os.path.dirname(__file__), "..", "..", "..", "exports")
    os.makedirs(export_dir, exist_ok=True)
    return os.path.join(export_dir, f"{slug}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx")