"""
Excel RAG metadata blueprint ingestion.

Reads the 13-sheet IP-SAKTI knowledge-engineering workbook (blueprint.xlsx) that
drives the RAG pipeline:

  Source Inventory      -> canonical source registry (list_reference_links etc.)
  RAG Metadata Template -> chunk metadata schema / defaults
  Evidence Passages     -> exact quoted legal text, converted into vector documents
  Rules                 -> deterministic legal rules (condition -> outcome)
  Sources / Source Versions -> authority-level registry (conflict resolution)

The workbook is optional. If no file is present every loader returns empty lists
and the RAG system operates purely on the curated text corpus.
"""

import os
import glob
import json
from typing import List, Dict, Any, Optional

BACKEND_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))


def blueprint_path() -> str:
    env_path = os.environ.get("IPSAKTI_BLUEPRINT_XLSX", "")
    if env_path and os.path.exists(os.path.abspath(env_path)):
        return os.path.abspath(env_path)
    data_dir = os.path.join(BACKEND_ROOT, "data")
    # Preferred: the knowledge-engineering workbook the user places in data/.
    for pattern in ("blueprint.xlsx", "IP_SAKTI_RAG_Metadata*.xlsx", "*_RAG_Metadata*.xlsx"):
        matches = sorted(glob.glob(os.path.join(data_dir, pattern)))
        if matches:
            return os.path.abspath(matches[0])
    return ""


def _normalize(value: Any) -> str:
    if value is None:
        return ""
    return " ".join(str(value).replace("_", " ").split()).strip().lower()


def _find_sheet(workbook, *names: str) -> Optional[Any]:
    targets = {_normalize(n) for n in names}
    for ws in workbook.worksheets:
        if _normalize(ws.title) in targets:
            return ws
    for ws in workbook.worksheets:
        if any(t in _normalize(ws.title) for t in targets):
            return ws
    return None


def _rows_to_dicts(ws) -> List[Dict[str, Any]]:
    rows = []
    headers = None
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        values = ["" if v is None else str(v).strip() for v in row]
        if not any(values):
            continue
        if headers is None:
            headers = values
            continue
        record = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if values[i]}
        if any(record.values()):
            rows.append(record)
    return rows


_HEADER_ALIASES = {
    "document id": "doc_id",
    "doc id": "doc_id",
    "title": "title",
    "authority": "authority",
    "source authority": "authority",
    "jurisdiction": "jurisdiction",
    "priority": "priority",
    "version": "version",
    "publication date": "publication_date",
    "effective date": "effective_date",
    "effective": "effective_date",
    "official url": "source_url",
    "url": "source_url",
    "ocr status": "ocr_status",
    "chunk status": "chunk_status",
    "embedding ready": "embedding_ready",
    "passage id": "passage_id",
    "passage_id": "passage_id",
    "section": "section_heading",
    "section reference": "section_reference",
    "exact text": "exact_passage",
    "exact passage": "exact_passage",
    "quoted text": "exact_passage",
    "act title": "act_title",
    "act": "act_title",
    "rule": "rule_name",
    "condition": "condition",
    "outcome": "outcome",
    "level": "authority_level",
    "authority level": "authority_level",
}


def _renamed(record: Dict[str, str]) -> Dict[str, str]:
    out: Dict[str, str] = {}
    for header, value in record.items():
        key = _HEADER_ALIASES.get(_normalize(header), _normalize(header))
        out[key] = value
    return out


def _g(record: Dict[str, Any], *names: str) -> str:
    """Tolerant lookup across key spellings (raw, normalized, aliased)."""
    if not record:
        return ""
    for name in names:
        variants = {name, _normalize(name), name.replace(" ", "_"), name.replace("_", " ")}
        for v in variants:
            if v in record and record[v]:
                return str(record[v])
            for k in record:
                if k.lower() == v:
                    val = record[k]
                    if val:
                        return str(val)
    return ""


def _persist_json(name: str, data: Any) -> None:
    knowledge_dir = os.path.join(BACKEND_ROOT, "app", "knowledge")
    os.makedirs(knowledge_dir, exist_ok=True)
    with open(os.path.join(knowledge_dir, name), "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)


def load_blueprint_documents() -> List[Dict[str, Any]]:
    """Read evidence passages + source inventory from the blueprint and return
    vector-ready documents with full RAG metadata (doc_id, authority, URL...)."""
    path = blueprint_path()
    if not path:
        return []
    try:
        import openpyxl
    except ImportError:
        return []

    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    documents: List[Dict[str, Any]] = []
    sources: List[Dict[str, Any]] = []
    rules: List[Dict[str, Any]] = []
    defaults: Dict[str, str] = {}

    ws_inventory = _find_sheet(wb, "source inventory", "sources", "folder checklist")
    if ws_inventory:
        for row in _rows_to_dicts(ws_inventory):
            r = _renamed(row)
            if r.get("doc_id") or r.get("title"):
                sources.append(r)

    ws_meta = _find_sheet(wb, "rag metadata template")
    if ws_meta:
        defaults = {_normalize(k): v for k, v in _rows_to_dicts(ws_meta)[0].items()} if _rows_to_dicts(ws_meta) else {}

    ws_evidence = _find_sheet(wb, "evidence passages")
    if ws_evidence:
        idx = 0
        for row in _rows_to_dicts(ws_evidence):
            r = _renamed(row)
            text = r.get("exact_passage", "").strip()
            if not text:
                continue
            doc_id = r.get("passage_id") or r.get("doc_id") or f"EVID-C{idx:03d}"
            authority = r.get("authority") or defaults.get("authority") or "Official Source"
            jurisdiction = r.get("jurisdiction") or defaults.get("jurisdiction") or ""
            title = r.get("act_title") or r.get("title") or defaults.get("title") or "Statutory Evidence Passage"
            documents.append({
                "content": text,
                "source": f"blueprint:{doc_id}",
                "category": "blueprint_evidence",
                "chunk_index": idx,
                "doc_id": doc_id,
                "title": title,
                "authority": authority,
                "jurisdiction": jurisdiction,
                "authority_level": int(r.get("authority_level") or defaults.get("authority_level") or 1),
                "source_url": r.get("source_url") or "",
                "effective_date": r.get("effective_date") or r.get("publication_date") or defaults.get("effective_date") or "",
                "section_heading": r.get("section_heading") or "",
                "citation_source": authority,
            })
            idx += 1

    # RAG Metadata Template: chunk-level summary rows -> RAG documents.
    _idx = 0
    ws_meta = _find_sheet(wb, "rag metadata template")
    if ws_meta:
        meta_rows = _rows_to_dicts(ws_meta)
        for row in meta_rows:
            r = _renamed(row)
            summary = _g(r, "document_summary", "document summary", "content")
            if not summary or summary.lower() in ("<example>", "n/a", "na", "--", "sample"):
                continue
            title = _g(r, "title") or defaults.get("title") or "RAG Metadata Record"
            doc_id = _g(r, "chunk_id", "chunk id", "doc_id", "document_id", "document id") or f"META-C{_idx:03d}"
            documents.append({
                "content": summary,
                "source": f"blueprint:{doc_id}",
                "category": "blueprint_chunk",
                "chunk_index": _idx,
                "doc_id": doc_id,
                "title": title,
                "authority": _g(r, "authority") or defaults.get("authority") or "Official Source",
                "jurisdiction": _g(r, "jurisdiction") or defaults.get("jurisdiction") or "",
                "authority_level": int(_g(r, "authority_level", "level") or defaults.get("authority_level") or 2),
                "source_url": _g(r, "source_url", "official url") or "",
                "effective_date": _g(r, "effective_date", "effective date", "publication_date") or "",
                "section_heading": _g(r, "section_heading", "section", "topic") or "",
                "version": _g(r, "version") or "",
                "document_type": _g(r, "document_type", "document type") or "",
                "language": _g(r, "language") or "",
                "citation_source": _g(r, "authority") or defaults.get("authority") or "Official Source",
            })
            _idx += 1

    # Source Inventory / Sources: catalog rows -> registry reference documents.
    _idx2 = 0
    ws_inv = _find_sheet(wb, "source inventory", "sources")
    if ws_inv:
        for row in _rows_to_dicts(ws_inv):
            r = _renamed(row)
            title = _g(r, "title", "document / webpage") or ""
            if not title or title.lower() in ("<example>", "n/a", "source"):
                continue
            desc = _g(r, "what to collect / why it matters",
                      "document_summary", "document summary", "recommended use") or ""
            authority = _g(r, "authority", "official authority") or "Official Source"
            jurisdiction = _g(r, "jurisdiction", "jurisdiction_hint") or ""
            url = _g(r, "source_url", "official url", "url", "canonical_url") or ""
            doc_id = _g(r, "doc_id", "document_id", "document id") or f"SRC-C{_idx2:03d}"
            content = f"{title}. {desc}" if desc else title
            documents.append({
                "content": content[:2000],
                "source": f"blueprint:{doc_id}",
                "category": "source_inventory",
                "chunk_index": _idx2,
                "doc_id": doc_id,
                "title": title,
                "authority": authority,
                "jurisdiction": jurisdiction,
                "authority_level": int(_g(r, "authority_level", "level") or 2),
                "source_url": url,
                "effective_date": _g(r, "publication_date", "publication date") or "",
                "section_heading": _g(r, "topic") or _g(r, "folder") or "",
                "priority": _g(r, "priority") or "",
                "recommended_use": _g(r, "recommended use") or "",
                "citation_source": authority,
            })
            _idx2 += 1

    # Glossary: canonical identity rows -> reference documents (botanical/IP terms).
    _idx3 = 0
    ws_gl = _find_sheet(wb, "glossary")
    if ws_gl:
        for row in _rows_to_dicts(ws_gl):
            r = _renamed(row)
            term = _g(r, "original_term", "original term", "term")
            identity = _g(r, "canonical_identity", "canonical identity", "alias")
            if not term or term.lower() in ("<example>", "n/a"):
                continue
            notes = _g(r, "ambiguity_notes", "ambiguity notes", "notes")
            content = f"{term} ({_g(r, 'language') or 'Sanskrit/Hindi'})"
            if identity:
                content += f" = {identity}"
            if notes:
                content += f". Ambiguity note: {notes}"
            documents.append({
                "content": content[:1500],
                "source": f"blueprint:GLOSS-{_idx3:03d}",
                "category": "glossary",
                "chunk_index": _idx3,
                "doc_id": f"GLOSS-{_idx3:03d}",
                "title": f"Glossary: {term}",
                "authority": "IP-SAKTI Knowledge Glossary",
                "jurisdiction": "International",
                "authority_level": 2,
                "source_url": "",
                "effective_date": "",
                "section_heading": "Glossary",
                "citation_source": "IP-SAKTI Knowledge Glossary",
            })
            _idx3 += 1

    ws_rules = _find_sheet(wb, "rules")
    if ws_rules:
        rules = [{**{_normalize(k): v for k, v in r.items()}, **{"_renamed": _renamed(r)}} for r in _rows_to_dicts(ws_rules)]

    # Persist derived registries so the retrieval + rules engine share them.
    if sources:
        _persist_json("blueprint_sources.json", sources)
    if rules:
        _persist_json("blueprint_rules.json", [
            {
                "condition": r.get("_renamed", {}).get("condition") or r.get("condition", ""),
                "outcome": r.get("_renamed", {}).get("outcome") or r.get("outcome", ""),
                "rule_name": r.get("_renamed", {}).get("rule_name") or r.get("rule_name", ""),
                "authority": r.get("_renamed", {}).get("authority") or "",
                "jurisdiction": r.get("_renamed", {}).get("jurisdiction") or "",
            }
            for r in rules
            if r
        ])

    wb.close()
    return documents


def load_blueprint_registry() -> Dict[str, Any]:
    """Return blueprint-derived source registry + rules (for the retrieval engine)."""
    path = blueprint_path()
    if not path:
        return {"sources": [], "rules": []}
    try:
        import openpyxl
    except ImportError:
        return {"sources": [], "rules": []}

    wb = openpyxl.load_workbook(path, read_only=False, data_only=True)
    sources: List[Dict[str, str]] = []
    rules: List[Dict[str, str]] = []

    ws_inventory = _find_sheet(wb, "source inventory", "sources")
    if ws_inventory:
        for row in _rows_to_dicts(ws_inventory):
            r = _renamed(row)
            if r.get("doc_id") or r.get("title"):
                sources.append(r)

    ws_rules = _find_sheet(wb, "rules")
    if ws_rules:
        for row in _rows_to_dicts(ws_rules):
            r = _renamed(row)
            if r.get("condition") and r.get("outcome"):
                rules.append(r)

    wb.close()
    return {"sources": sources, "rules": rules}